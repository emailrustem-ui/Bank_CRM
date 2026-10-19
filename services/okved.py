import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from models import db, utcnow, Product, OkvedRule, OkvedCheck, ReferenceVersion, Company, CompanyOkved

CODE_RE = re.compile(r"\d{2}(?:\.\d{1,3}){0,3}")

STATUS_RANK = {"unknown": 0, "yes": 1, "mixed": 2, "no": 3}

STATUS_VOCAB = {
    "ДА", "НЕТ", "ДА/НЕТ", "НЕТ/ДА",
    "YES", "NO", "YES/NO", "NO/YES",
}


def normalize_code(value):
    if value is None:
        return ""
    text = str(value).strip().replace(",", ".")
    match = CODE_RE.search(text)
    if not match:
        return ""
    return match.group(0).rstrip(".")


def normalize_status(value):
    if value is None:
        return "unknown"

    text = str(value).strip().upper().replace("Ё", "Е")
    if not text:
        return "unknown"

    compact = re.sub(r"\s+", "", text)

    if compact in {"ДА", "YES", "Y", "+", "РАЗРЕШЕН"}:
        return "yes"
    if compact in {"НЕТ", "NO", "N", "-", "НЕ", "НЕРАЗРЕШЕН", "ЗАПРЕЩЕН"}:
        return "no"

    has_yes = "ДА" in text or "YES" in text
    has_no = "НЕТ" in text or "НЕ" in text or "ЗАПРЕТ" in text or "НЕРАЗРЕШ" in compact

    if has_yes and has_no:
        return "mixed"
    if has_no:
        return "no"
    if has_yes or "РАЗРЕШ" in text:
        return "yes"
    if "УСЛОВ" in text:
        return "mixed"

    return "unknown"


def _looks_status(value):
    text = str(value or "").strip().upper().replace("Ё", "Е")
    return text in STATUS_VOCAB


def detect_scope(condition):
    if not condition:
        return "any"

    text = condition.lower()

    if re.search(r"только[^.]{0,60}дополнительн", text):
        return "additional"

    if (
        "дополнительн" in text
        and "основн" in text
        and re.search(r"основн[^.]{0,80}(невозможен|нельзя|запрещ|невозможно|исключ)", text)
    ):
        return "additional"

    if re.search(r"только[^.]{0,60}основн", text):
        return "main"

    return "any"


def _is_ole2(file_path):
    """Старый Excel 97-2003 (.xls) начинается с сигнатуры D0 CF 11 E0."""
    try:
        with open(file_path, "rb") as fh:
            return fh.read(4) == b"\xd0\xcf\x11\xe0"
    except Exception:
        return False


def _convert_xls_to_xlsx(xls_path, xlsx_path):
    import xlrd

    xlrd_book = xlrd.open_workbook(str(xls_path))
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name in xlrd_book.sheet_names():
        xlrd_sheet = xlrd_book.sheet_by_name(sheet_name)
        ws = wb.create_sheet(title=sheet_name)

        for row_idx in range(xlrd_sheet.nrows):
            for col_idx in range(xlrd_sheet.ncols):
                cell = xlrd_sheet.cell(row_idx, col_idx)
                ws.cell(row=row_idx + 1, column=col_idx + 1, value=cell.value)

    wb.save(str(xlsx_path))


def _merged_value(ws, row, col, merged_ranges):
    for rng in merged_ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value


def _detect_header(ws):
    max_row = min(ws.max_row or 1, 40)
    max_col = ws.max_column or 1

    for row_idx in range(1, max_row + 1):
        values = []
        for col_idx in range(1, max_col + 1):
            value = ws.cell(row_idx, col_idx).value
            values.append(str(value or "").strip().lower())

        code_idx = None
        for idx, value in enumerate(values):
            if "оквэд" in value or "оквед" in value or value in {"код", "код вида деятельности"}:
                code_idx = idx
                break

        if code_idx is None:
            continue

        condition_idx = None
        for idx, value in enumerate(values):
            if idx == code_idx:
                continue
            if any(word in value for word in ("услов", "примеч", "коммент", "поясн")):
                condition_idx = idx
                break

        return {
            "row": row_idx,
            "code": code_idx,
            "condition": condition_idx,
        }

    return None


def _detect_status_columns(ws, header_info, merged_ranges):
    max_col = ws.max_column or 1
    header_row = header_info["row"]
    code_idx = header_info["code"]

    def _row_values(row_idx):
        return [
            str(ws.cell(row_idx, c).value or "").strip().lower()
            for c in range(1, max_col + 1)
        ]

    cols = []
    for idx, value in enumerate(_row_values(header_row)):
        if idx == code_idx:
            continue
        if "разреш" in value or "статус" in value or "допустим" in value:
            cols.append(idx)

    if not cols and header_row + 1 <= (ws.max_row or 1):
        for idx, value in enumerate(_row_values(header_row + 1)):
            if idx == code_idx:
                continue
            if "разреш" in value or "статус" in value or "допустим" in value:
                cols.append(idx)

    if not cols:
        for col_idx in range(1, max_col + 1):
            idx = col_idx - 1
            if idx == code_idx:
                continue
            hits = 0
            total = 0
            for row_idx in range(header_row + 1, min(header_row + 41, (ws.max_row or 1) + 1)):
                value = _merged_value(ws, row_idx, col_idx, merged_ranges)
                if value is None or not str(value).strip():
                    continue
                total += 1
                if _looks_status(value):
                    hits += 1
            if total >= 3 and hits / total >= 0.6:
                cols.append(idx)
            if len(cols) == 2:
                break

    return cols[:2]


def _guess_condition_column(ws, header_info, merged_ranges, status_cols):
    if header_info["condition"] is not None:
        return header_info["condition"]

    start = header_info["row"] + 1
    max_col = ws.max_column or 1
    blocked = set(status_cols) | {header_info["code"]}

    for col_idx in range(1, max_col + 1):
        idx = col_idx - 1
        if idx in blocked:
            continue
        for row_idx in range(start, min(start + 40, (ws.max_row or 1) + 1)):
            value = _merged_value(ws, row_idx, col_idx, merged_ranges)
            if value and len(str(value).strip()) > 70:
                return idx

    return None


def _sheet_intro(ws, header_row):
    lines = []
    for row_idx in range(1, header_row):
        for col_idx in range(1, (ws.max_column or 1) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is None:
                continue
            text = str(value).strip()
            if len(text) > 45:
                lines.append(text)
    return "\n".join(lines)[:4000]


def _product_name_from_sheet(ws, sheet_name, header_row):
    name = (sheet_name or "").strip()
    if name and name.lower() not in {"лист1", "лист 1", "sheet1", "sheet 1"}:
        return name[:300]

    for row_idx in range(1, header_row):
        for col_idx in range(1, (ws.max_column or 1) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value:
                text = str(value).strip()
                if 5 < len(text) < 150:
                    return text[:300]

    return name or "Продукт"


def _combine_statuses(status_main, status_alt):
    if status_alt == "unknown":
        return status_main, None

    if status_main == status_alt:
        return status_main, "any"
    if status_main == "yes" and status_alt == "no":
        return "yes", "main"
    if status_main == "no" and status_alt == "yes":
        return "yes", "additional"
    if status_main == "unknown":
        return status_alt, "any"
    return "mixed", "any"


def import_reference(file_path, filename):
    file_path = Path(file_path)

    # Определяем реальный формат по содержимому, а не по расширению
    temp_path = None
    if _is_ole2(file_path):
        temp_path = file_path.with_suffix(file_path.suffix + ".temp.xlsx")
        _convert_xls_to_xlsx(file_path, temp_path)
        file_path = temp_path

    try:
        workbook = load_workbook(str(file_path), data_only=True)

        version = ReferenceVersion(
            filename=filename,
            created_at=utcnow(),
            is_active=True,
        )
        db.session.add(version)
        db.session.flush()

        for old in ReferenceVersion.query.filter(ReferenceVersion.id != version.id).all():
            old.is_active = False

        merged_rules = {}
        product_descriptions = {}
        skipped_sheets = []

        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            header_info = _detect_header(ws)

            if header_info is None:
                skipped_sheets.append(sheet_name)
                continue

            product_name = _product_name_from_sheet(ws, sheet_name, header_info["row"])
            intro = _sheet_intro(ws, header_info["row"])
            product_descriptions[product_name] = intro

            merged_ranges = list(ws.merged_cells.ranges)
            status_cols = _detect_status_columns(ws, header_info, merged_ranges)
            condition_col = _guess_condition_column(ws, header_info, merged_ranges, status_cols)

            if not status_cols:
                skipped_sheets.append(sheet_name)
                continue

            code_col = header_info["code"] + 1
            status_col = status_cols[0] + 1
            status2_col = (status_cols[1] + 1) if len(status_cols) > 1 else None

            last_condition = ""

            for row_idx in range(header_info["row"] + 1, (ws.max_row or 1) + 1):
                raw_code = _merged_value(ws, row_idx, code_col, merged_ranges)
                code = normalize_code(raw_code)

                condition_value = ""
                if condition_col is not None:
                    value = _merged_value(ws, row_idx, condition_col + 1, merged_ranges)
                    if value is not None and str(value).strip():
                        condition_value = str(value).strip()
                        last_condition = condition_value

                if not code:
                    continue

                raw_status = str(_merged_value(ws, row_idx, status_col, merged_ranges) or "").strip()
                raw_status2 = ""
                if status2_col is not None:
                    raw_status2 = str(_merged_value(ws, row_idx, status2_col, merged_ranges) or "").strip()

                status_main = normalize_status(raw_status)
                status_alt = normalize_status(raw_status2) if raw_status2 else "unknown"

                status_norm, scope = _combine_statuses(status_main, status_alt)

                effective_condition = condition_value or last_condition

                if scope is None:
                    scope = detect_scope(effective_condition)

                if status_norm == "unknown" and effective_condition and len(effective_condition) > 40:
                    status_norm = "mixed"

                key = (product_name, code)

                if key not in merged_rules:
                    merged_rules[key] = {
                        "product_name": product_name,
                        "code": code,
                        "status_raw": raw_status + (f" / {raw_status2}" if raw_status2 else ""),
                        "status_norm": status_norm,
                        "scope": scope,
                        "condition": effective_condition,
                    }
                    continue

                existing = merged_rules[key]

                if STATUS_RANK.get(status_norm, 0) > STATUS_RANK.get(existing["status_norm"], 0):
                    existing["status_norm"] = status_norm
                    existing["status_raw"] = raw_status
                    existing["scope"] = scope

                if effective_condition and effective_condition not in existing["condition"]:
                    existing["condition"] = (existing["condition"] + "\n" + effective_condition).strip()

        product_map = {}

        for (product_name, code), rule_data in merged_rules.items():
            if product_name not in product_map:
                product = Product.query.filter_by(name=product_name).first()
                if not product:
                    product = Product(name=product_name, is_active=True)
                    db.session.add(product)
                    db.session.flush()
                else:
                    product.is_active = True

                if product_descriptions.get(product_name) and not product.description:
                    product.description = product_descriptions[product_name]

                product_map[product_name] = product

            product = product_map[product_name]

            rule = OkvedRule(
                version_id=version.id,
                product_id=product.id,
                code=code,
                status_raw=rule_data["status_raw"],
                status_norm=rule_data["status_norm"],
                condition=rule_data["condition"],
                scope=rule_data["scope"],
            )
            db.session.add(rule)

        version.product_count = len(product_map)
        version.rule_count = len(merged_rules)

        db.session.commit()

        return {
            "version_id": version.id,
            "products": len(product_map),
            "rules": len(merged_rules),
            "skipped": skipped_sheets,
        }

    finally:
        if temp_path:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:
                pass


# ============================================================
# ПРОВЕРКА КОМПАНИИ: ПРИОРИТЕТ ОСНОВНОГО ОКВЭД
# ============================================================

def _find_rule(code, rules_by_code):
    if not code:
        return None

    if code in rules_by_code:
        return rules_by_code[code]

    parts = code.split(".")
    for length in range(len(parts) - 1, 0, -1):
        parent = ".".join(parts[:length])
        if parent in rules_by_code:
            return rules_by_code[parent]

    return None


def _status_word(norm):
    return {
        "yes": "разрешён",
        "no": "запрещён",
        "mixed": "условно",
    }.get(norm, "статус не распознан")


def evaluate_product(company_okveds, rules_by_code, product_description=""):
    primary = None
    additions = []

    for item in company_okveds:
        code = normalize_code(item.code)
        if not code:
            continue
        if item.is_main and primary is None:
            primary = type("_", (), {"code": code, "is_main": True})()
        else:
            additions.append(type("_", (), {"code": code, "is_main": False})())

    if primary is None and company_okveds:
        primary = type("_", (), {"code": normalize_code(company_okveds[0].code), "is_main": True})()

    def _block(rule, kind):
        text = f"{rule.code} ({kind}) — {_status_word(rule.status_norm)}."
        if rule.condition:
            text += f"\n{rule.condition}"
        return text

    rule_primary = _find_rule(primary.code, rules_by_code) if primary else None

    if rule_primary and rule_primary.scope == "additional":
        detail = (
            f"{rule_primary.code} (основной) — НЕДОСТУПНО: "
            f"продукт возможен только при наличии ДОПОЛНИТЕЛЬНОГО ОКВЭД."
        )
        if rule_primary.condition:
            detail += f"\n{rule_primary.condition}"
        return "forbidden", detail

    if rule_primary:
        if rule_primary.status_norm == "yes":
            lines = [_block(rule_primary, "основной")]

            for add in additions:
                rule_add = _find_rule(add.code, rules_by_code)
                if rule_add and rule_add.status_norm == "no" and rule_add.scope in ("any", "additional"):
                    lines.append(
                        f"Внимание: дополнительный ОКВЭД {rule_add.code} по этому продукту запрещён."
                        + (f"\n{rule_add.condition}" if rule_add.condition else "")
                    )

            detail = "\n\n".join(lines)
            if product_description:
                detail += f"\n\nОбщие условия продукта:\n{product_description}"
            return "allowed", detail

        if rule_primary.status_norm == "no":
            return "forbidden", _block(rule_primary, "основной")

        detail = _block(rule_primary, "основной")
        if product_description:
            detail += f"\n\nОбщие условия продукта:\n{product_description}"
        return "conditional", detail

    forbidden, conditional, allowed = [], [], []

    for add in additions:
        rule_add = _find_rule(add.code, rules_by_code)
        if not rule_add or rule_add.scope == "main":
            continue

        if rule_add.status_norm == "no":
            forbidden.append(_block(rule_add, "дополнительный"))
        elif rule_add.status_norm == "mixed":
            conditional.append(_block(rule_add, "дополнительный"))
        elif rule_add.status_norm == "yes":
            allowed.append(_block(rule_add, "дополнительный"))

    if forbidden:
        return "forbidden", "\n\n".join(forbidden)
    if conditional:
        return "conditional", "\n\n".join(conditional)
    if allowed:
        detail = "\n\n".join(allowed)
        if product_description:
            detail += f"\n\nОбщие условия продукта:\n{product_description}"
        return "allowed", detail

    if primary:
        return "unknown", (
            f"Основной ОКВЭД {primary.code} не найден в справочнике продукта. "
            f"Дополнительные ОКВЭД соответствия не дали."
        )

    return "unknown", "У компании не заполнен основной ОКВЭД."


def check_company(company):
    version = (
        ReferenceVersion.query.filter_by(is_active=True)
        .order_by(ReferenceVersion.id.desc())
        .first()
    )

    if not version:
        return []

    okveds = list(company.okveds)

    if not okveds and company.okved_main_code:
        item = CompanyOkved(
            company_id=company.id,
            code=normalize_code(company.okved_main_code),
            title=company.okved_main_title or "",
            is_main=True,
        )
        db.session.add(item)
        db.session.flush()
        okveds = [item]

    products = Product.query.filter_by(is_active=True).order_by(Product.name).all()
    results = []

    for product in products:
        rules = OkvedRule.query.filter_by(
            version_id=version.id,
            product_id=product.id,
        ).all()

        rules_by_code = {rule.code: rule for rule in rules}

        result, detail = evaluate_product(okveds, rules_by_code, product.description or "")

        check = OkvedCheck.query.filter_by(
            company_id=company.id,
            product_id=product.id,
        ).first()

        if not check:
            check = OkvedCheck(company_id=company.id, product_id=product.id)
            db.session.add(check)

        check.result = result
        check.detail = detail
        check.version_id = version.id
        check.checked_at = utcnow()

        results.append(check)

    db.session.flush()
    return results


def recheck_all_companies():
    companies = Company.query.all()
    for company in companies:
        check_company(company)
    db.session.commit()
    return len(companies)