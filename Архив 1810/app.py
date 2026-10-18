import io
import json
import re
import threading
import time
import uuid
import webbrowser
from datetime import date, datetime, timedelta
from pathlib import Path

from flask import (
    Flask, abort, flash, jsonify, redirect,
    render_template, request, send_file, url_for,
)
from flask_login import (
    LoginManager, current_user, login_required,
    login_user, logout_user,
)
from flask_wtf.csrf import CSRFProtect
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.utils import secure_filename

from config import Config, DATA_DIR, REFERENCES_DIR
from models import (
    ACTIVITY_TYPES, COMPANY_STATUSES, CHECK_RESULTS, DEAL_STATUSES,
    DEFAULT_TEMPLATE_BODY, DEFAULT_TEMPLATE_SUBJECT,
    Activity, Company, CompanyOkved, Contact, Deal, DealStatusLog,
    MailTemplate, OkvedCheck, OkvedRule, Product, ReferenceVersion,
    Task, User, db,
)
from services import okved as okved_service
from services.checko import CheckoError, apply_scraped_data, scrape_company
from services.mail import (
    build_greeting, mailto_link, prepare_email, try_open_outlook,
)

for folder in (DATA_DIR, REFERENCES_DIR, DATA_DIR / "snapshots", DATA_DIR / "chrome_profile"):
    folder.mkdir(parents=True, exist_ok=True)


app = Flask(__name__)
app.config.from_object(Config)

db.init_app(app)
csrf = CSRFProtect(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"
login_manager.login_message = "Войдите в систему."
login_manager.login_message_category = "warning"


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


@app.context_processor
def inject_globals():
    return {
        "company_statuses": COMPANY_STATUSES,
        "activity_types": ACTIVITY_TYPES,
        "check_results": CHECK_RESULTS,
        "deal_statuses": DEAL_STATUSES,
        "now": datetime.utcnow(),
    }


@app.template_filter("money")
def money_filter(value):
    try:
        amount = float(value or 0)
    except (TypeError, ValueError):
        return "—"
    if not amount:
        return "—"
    return f"{amount:,.0f}".replace(",", " ") + " руб."


def _parse_amount(raw):
    text = str(raw or "").strip().lower()
    if not text:
        return 0.0

    multiplier = 1.0
    if "млрд" in text:
        multiplier = 1_000_000_000.0
        text = text.replace("млрд", "")
    elif "млн" in text:
        multiplier = 1_000_000.0
        text = text.replace("млн", "")
    elif "тыс" in text:
        multiplier = 1_000.0
        text = text.replace("тыс", "")

    text = text.replace("руб", "").replace(" ", "").replace(",", ".")
    text = re.sub(r"[^\d.]", "", text)

    try:
        return float(text) * multiplier
    except ValueError:
        return 0.0


def _log_deal_status(deal, old_status, new_status):
    if old_status == new_status:
        return
    db.session.add(DealStatusLog(
        deal_id=deal.id,
        company_id=deal.company_id,
        old_status=old_status,
        new_status=new_status,
        user_id=current_user.id if current_user.is_authenticated else None,
    ))


# ============================================================
# ФОНОВЫЕ ЗАДАЧИ ЗАГРУЗКИ ДАННЫХ
# ============================================================

JOBS = {}
JOBS_LOCK = threading.Lock()


def _job_update(job_id, **kwargs):
    with JOBS_LOCK:
        job = JOBS.setdefault(job_id, {})
        job.update(kwargs)
        job["updated_at"] = datetime.utcnow()


def _job_get(job_id):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        return dict(job) if job else None


def _fetch_worker(job_id, inn, user_id, company_id=None):
    with app.app_context():
        try:
            _job_update(job_id, status="running", progress=10, message="Запускаю загрузку данных…")

            data = scrape_company(
                inn,
                log=lambda msg: _job_update(job_id, message=msg)
            )

            _job_update(job_id, progress=65, message="Сохраняю карточку компании…")

            if company_id:
                company = db.session.get(Company, company_id)
            else:
                company = Company.query.filter_by(inn=inn).first()

            if not company:
                company = Company(
                    inn=inn,
                    name=f"Компания {inn}",
                    manager_id=user_id,
                    status="new",
                )
                db.session.add(company)
                db.session.flush()

            apply_scraped_data(company, data)

            # Нормализация телефонов: разделитель только ";",
            # пробелы внутри номера сохраняются (+7 347 270-13-46)
            raw_phones = str(data.get("phone") or "")
            phone_parts = [p.strip() for p in raw_phones.split(";") if p.strip()]
            if phone_parts:
                company.phone = "; ".join(phone_parts)
                for contact in company.contacts:
                    contact.phone = "; ".join(phone_parts)

            db.session.commit()

            _job_update(job_id, progress=85, message="Проверяю ОКВЭД по продуктам…")

            checks = okved_service.check_company(company)
            db.session.commit()

            db.session.add(Activity(
                company_id=company.id,
                user_id=user_id,
                type="note",
                subject="Загрузка данных компании",
                body=(
                    f"Получены данные по ИНН {inn}. "
                    f"ОКВЭД: {len(company.okveds)} шт. "
                    f"Проверено продуктов: {len(checks)}."
                ),
            ))
            db.session.commit()

            _job_update(
                job_id,
                status="done",
                progress=100,
                message="Готово. Компания создана и проверена по продуктам.",
                company_id=company.id,
            )

        except Exception as exc:
            db.session.rollback()
            _job_update(
                job_id,
                status="error",
                progress=100,
                message="Не удалось загрузить компанию.",
                error=str(exc),
            )


# ============================================================
# ИНИЦИАЛИЗАЦИЯ
# ============================================================

def _ensure_company_extra_column():
    from sqlalchemy import text as sa_text

    try:
        rows = db.session.execute(sa_text("PRAGMA table_info(companies)")).fetchall()
        columns = {row[1] for row in rows}
        if columns and "extra" not in columns:
            db.session.execute(
                sa_text("ALTER TABLE companies ADD COLUMN extra TEXT DEFAULT ''")
            )
            db.session.commit()
    except Exception:
        db.session.rollback()


def init_database():
    with app.app_context():
        db.create_all()
        _ensure_company_extra_column()

        if not User.query.first():
            user = User(username="manager", full_name="Клиентский менеджер", is_enabled=True)
            user.set_password("manager")
            db.session.add(user)

        if not MailTemplate.query.first():
            db.session.add(MailTemplate(
                name="Первичное обращение",
                subject=DEFAULT_TEMPLATE_SUBJECT,
                body=DEFAULT_TEMPLATE_BODY,
                is_default=True,
            ))

        db.session.commit()


# ============================================================
# АВТОРИЗАЦИЯ И ЛИЧНЫЙ КАБИНЕТ
# ============================================================

@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        user = User.query.filter_by(username=username).first()
        if user and user.is_enabled and user.check_password(password):
            login_user(user)
            return redirect(url_for("dashboard"))

        flash("Неверный логин или пароль.", "danger")

    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    if request.method == "POST":
        action = request.form.get("action", "")

        if action == "save_profile":
            current_user.full_name = request.form.get("full_name", "").strip()[:200]
            db.session.commit()
            flash("Профиль сохранён. Имя используется в подписи писем.", "success")

        elif action == "change_password":
            current_password = request.form.get("current_password", "")
            new_password = request.form.get("new_password", "")
            confirm_password = request.form.get("confirm_password", "")

            if not current_user.check_password(current_password):
                flash("Текущий пароль указан неверно.", "danger")
            elif len(new_password) < 6:
                flash("Новый пароль должен быть не короче 6 символов.", "danger")
            elif new_password != confirm_password:
                flash("Новый пароль и подтверждение не совпадают.", "danger")
            else:
                current_user.set_password(new_password)
                db.session.commit()
                flash("Пароль изменён.", "success")

        return redirect(url_for("profile"))

    return render_template("profile.html")


# ============================================================
# ДАШБОРД
# ============================================================

@app.route("/")
@login_required
def dashboard():
    companies_count = Company.query.count()
    checks_count = OkvedCheck.query.count()
    allowed_count = OkvedCheck.query.filter_by(result="allowed").count()
    products_count = Product.query.filter_by(is_active=True).count()

    open_tasks = (
        Task.query.filter_by(status="open")
        .order_by(Task.due_date.is_(None), Task.due_date)
        .limit(10)
        .all()
    )

    overdue_count = Task.query.filter(
        Task.status == "open",
        Task.due_date.isnot(None),
        Task.due_date < date.today(),
    ).count()

    recent_companies = (
        Company.query.order_by(Company.created_at.desc()).limit(6).all()
    )

    recent_activities = (
        Activity.query.order_by(Activity.created_at.desc()).limit(6).all()
    )

    funnel = []
    for key, label in DEAL_STATUSES.items():
        deals = Deal.query.filter_by(status=key).all()
        funnel.append({
            "key": key,
            "label": label,
            "count": len(deals),
            "amount": sum(d.amount or 0 for d in deals),
        })

    soon = date.today() + timedelta(days=7)
    reminders = (
        Deal.query.filter(
            Deal.status.notin_(["issued", "rejected"]),
            Deal.term_date.isnot(None),
            Deal.term_date <= soon,
        )
        .order_by(Deal.term_date)
        .all()
    )

    status_logs = (
        DealStatusLog.query.order_by(DealStatusLog.created_at.desc()).limit(8).all()
    )

    active_deals = Deal.query.filter(
        Deal.status.notin_(["issued", "rejected"])
    ).count()

    return render_template(
        "dashboard.html",
        companies_count=companies_count,
        checks_count=checks_count,
        allowed_count=allowed_count,
        products_count=products_count,
        open_tasks=open_tasks,
        overdue_count=overdue_count,
        recent_companies=recent_companies,
        recent_activities=recent_activities,
        active_deals=active_deals,
        funnel=funnel,
        reminders=reminders,
        status_logs=status_logs,
    )


# ============================================================
# КОМПАНИИ
# ============================================================

@app.route("/companies")
@login_required
def companies():
    query = Company.query

    search = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()

    if search:
        like = f"%{search}%"
        query = query.filter(
            db.or_(Company.inn.ilike(like), Company.name.ilike(like), Company.short_name.ilike(like))
        )

    if status:
        query = query.filter_by(status=status)

    page = request.args.get("page", 1, type=int)
    pagination = (
        query.order_by(Company.created_at.desc())
        .paginate(page=page, per_page=Config.COMPANIES_PER_PAGE, error_out=False)
    )

    return render_template(
        "companies_list.html",
        pagination=pagination,
        companies=pagination.items,
        search=search,
        status=status,
    )


@app.route("/companies/new", methods=["GET", "POST"])
@login_required
def company_new():
    job_id = request.args.get("job")
    if job_id:
        job = _job_get(job_id)
        if not job:
            flash("Задача не найдена. Возможно, сервер был перезапущен.", "warning")
            return render_template("company_new.html", inn="", error=None)
        return render_template("company_fetch.html", job_id=job_id, job=job)

    if request.method == "POST":
        mode = request.form.get("mode", "auto")
        inn = re.sub(r"\D", "", request.form.get("inn", ""))

        if mode == "manual":
            if not inn:
                flash("Для ручной карточки тоже нужен ИНН — это уникальный ключ компании.", "danger")
                return render_template("company_new.html", inn=inn, error=None)

            existing = Company.query.filter_by(inn=inn).first()
            if existing:
                flash("Такая компания уже есть в CRM.", "warning")
                return redirect(url_for("company_detail", company_id=existing.id))

            company = Company(inn=inn, name=f"Компания {inn}", manager_id=current_user.id)
            db.session.add(company)
            db.session.commit()
            okved_service.check_company(company)
            db.session.commit()

            flash("Карточка создана. Заполните реквизиты вручную.", "success")
            return redirect(url_for("company_detail", company_id=company.id))

        if len(inn) not in (10, 12):
            flash("ИНН должен содержать 10 или 12 цифр.", "danger")
            return render_template("company_new.html", inn=inn, error=None)

        existing = Company.query.filter_by(inn=inn).first()
        if existing:
            flash("Компания с таким ИНН уже есть в CRM.", "warning")
            return redirect(url_for("company_detail", company_id=existing.id))

        new_job_id = uuid.uuid4().hex
        _job_update(new_job_id, status="running", progress=0, message="Задача поставлена в очередь…")

        thread = threading.Thread(
            target=_fetch_worker,
            args=(new_job_id, inn, current_user.id),
            daemon=True,
        )
        thread.start()

        return redirect(url_for("company_new", job=new_job_id))

    return render_template("company_new.html", inn="", error=None)


@app.route("/company/<int:company_id>")
@login_required
def company_detail(company_id):
    company = db.get_or_404(Company, company_id)

    checks = (
        OkvedCheck.query.filter_by(company_id=company.id)
        .join(Product)
        .order_by(Product.name)
        .all()
    )

    activities = (
        Activity.query.filter_by(company_id=company.id)
        .order_by(Activity.created_at.desc())
        .all()
    )

    tasks = (
        Task.query.filter_by(company_id=company.id)
        .order_by(Task.status, Task.due_date)
        .all()
    )

    templates = MailTemplate.query.order_by(MailTemplate.is_default.desc(), MailTemplate.name).all()
    version = ReferenceVersion.query.filter_by(is_active=True).first()

    extra_data = {}
    try:
        extra_data = json.loads(company.extra or "{}") or {}
    except Exception:
        extra_data = {}

    facts = extra_data.get("facts", {}) or {}

    return render_template(
        "company_detail.html",
        company=company,
        checks=checks,
        activities=activities,
        tasks=tasks,
        templates=templates,
        active_version=version,
        extra=extra_data,
        facts=facts,
    )


@app.route("/company/<int:company_id>/edit", methods=["POST"])
@login_required
def company_edit(company_id):
    company = db.get_or_404(Company, company_id)

    fields = {
        "name": 500,
        "short_name": 500,
        "director": 300,
        "address": 100000,
        "phone": 300,
        "email": 300,
        "site": 300,
        "reg_date": 80,
        "ogrn": 20,
    }

    for field, limit in fields.items():
        if field in request.form:
            setattr(company, field, request.form.get(field, "").strip()[:limit])

    if request.form.get("status") in COMPANY_STATUSES:
        company.status = request.form["status"]

    new_main = okved_service.normalize_code(request.form.get("okved_main_code", ""))
    if new_main:
        company.okved_main_code = new_main
        main_item = CompanyOkved.query.filter_by(company_id=company.id, is_main=True).first()
        if main_item:
            main_item.code = new_main
        else:
            db.session.add(CompanyOkved(
                company_id=company.id, code=new_main, title=company.okved_main_title, is_main=True
            ))

    db.session.commit()
    okved_service.check_company(company)
    db.session.commit()

    flash("Карточка сохранена. Проверка ОКВЭД обновлена.", "success")
    return redirect(url_for("company_detail", company_id=company.id))


@app.route("/company/<int:company_id>/delete", methods=["POST"])
@login_required
def company_delete(company_id):
    company = db.get_or_404(Company, company_id)
    inn = company.inn
    name = company.display_name

    db.session.delete(company)
    db.session.commit()

    flash(f"Компания {name} (ИНН {inn}) удалена вместе со связанными данными.", "success")
    return redirect(url_for("companies"))


@app.route("/company/<int:company_id>/refresh", methods=["POST"])
@login_required
def company_refresh(company_id):
    company = db.get_or_404(Company, company_id)

    job_id = uuid.uuid4().hex
    _job_update(job_id, status="running", progress=0, message="Запускаю обновление…")

    thread = threading.Thread(
        target=_fetch_worker,
        args=(job_id, company.inn, current_user.id, company.id),
        daemon=True,
    )
    thread.start()

    return redirect(url_for("company_new", job=job_id))


@app.route("/company/<int:company_id>/check", methods=["POST"])
@login_required
def company_check(company_id):
    company = db.get_or_404(Company, company_id)
    checks = okved_service.check_company(company)
    db.session.commit()

    if not checks:
        flash("Справочник ОКВЭД не загружен. Загрузите файл в разделе «Продукты».", "warning")
    else:
        flash(f"Проверено продуктов: {len(checks)}.", "success")

    return redirect(url_for("company_detail", company_id=company.id, tab="okved"))


# ============================================================
# ОКВЭД КОМПАНИИ
# ============================================================

@app.route("/company/<int:company_id>/okved/add", methods=["POST"])
@login_required
def okved_add(company_id):
    company = db.get_or_404(Company, company_id)

    code = okved_service.normalize_code(request.form.get("code", ""))
    if not code:
        flash("Введите корректный код ОКВЭД, например 46.12.1.", "danger")
        return redirect(url_for("company_detail", company_id=company.id, tab="okved"))

    title = request.form.get("title", "").strip()
    is_main = request.form.get("is_main") == "on"

    existing = CompanyOkved.query.filter_by(company_id=company.id, code=code).first()
    if existing:
        flash(f"ОКВЭД {code} уже добавлен.", "warning")
        return redirect(url_for("company_detail", company_id=company.id, tab="okved"))

    if is_main:
        CompanyOkved.query.filter_by(company_id=company.id, is_main=True).update({"is_main": False})
        company.okved_main_code = code
        company.okved_main_title = title

    db.session.add(CompanyOkved(
        company_id=company.id, code=code, title=title, is_main=is_main
    ))
    db.session.commit()

    okved_service.check_company(company)
    db.session.commit()

    flash(f"ОКВЭД {code} добавлен, проверка обновлена.", "success")
    return redirect(url_for("company_detail", company_id=company.id, tab="okved"))


@app.route("/company/<int:company_id>/okved/<int:okved_id>/delete", methods=["POST"])
@login_required
def okved_delete(company_id, okved_id):
    company = db.get_or_404(Company, company_id)
    item = db.get_or_404(CompanyOkved, okved_id)

    if item.company_id != company.id:
        abort(404)

    was_main = item.is_main
    code = item.code

    db.session.delete(item)
    db.session.flush()

    if was_main:
        company.okved_main_code = ""
        company.okved_main_title = ""

    db.session.commit()
    okved_service.check_company(company)
    db.session.commit()

    flash(f"ОКВЭД {code} удалён, проверка обновлена.", "success")
    return redirect(url_for("company_detail", company_id=company.id, tab="okved"))


# ============================================================
# КОНТАКТЫ
# ============================================================

@app.route("/company/<int:company_id>/contact/save", methods=["POST"])
@login_required
def contact_save(company_id):
    company = db.get_or_404(Company, company_id)
    contact_id = request.form.get("contact_id", type=int)

    if contact_id:
        contact = db.get_or_404(Contact, contact_id)
        if contact.company_id != company.id:
            abort(404)
    else:
        contact = Contact(company_id=company.id)
        db.session.add(contact)

    contact.full_name = request.form.get("full_name", "").strip()[:300]
    contact.position = request.form.get("position", "").strip()[:200]
    contact.phone = request.form.get("phone", "").strip()[:300]
    contact.email = request.form.get("email", "").strip()[:300]
    contact.is_primary = request.form.get("is_primary") == "on"

    if not contact.full_name and not contact.email:
        flash("Заполните ФИО или e-mail контакта.", "danger")
        db.session.rollback()
        return redirect(url_for("company_detail", company_id=company.id, tab="contacts"))

    if contact.is_primary:
        Contact.query.filter_by(company_id=company.id, is_primary=True).update({"is_primary": False})
        contact.is_primary = True

    db.session.commit()

    flash("Контакт сохранён.", "success")
    return redirect(url_for("company_detail", company_id=company.id, tab="contacts"))


@app.route("/company/<int:company_id>/contact/<int:contact_id>/delete", methods=["POST"])
@login_required
def contact_delete(company_id, contact_id):
    contact = db.get_or_404(Contact, contact_id)
    if contact.company_id != company_id:
        abort(404)

    db.session.delete(contact)
    db.session.commit()

    flash("Контакт удалён.", "success")
    return redirect(url_for("company_detail", company_id=company_id, tab="contacts"))


# ============================================================
# ЗАЯВКИ
# ============================================================

def _fill_deal(deal, form):
    company_id = form.get("company_id", type=int)
    if company_id:
        deal.company_id = company_id

    product_id = form.get("product_id", type=int)
    deal.product_id = product_id or None

    deal.amount = _parse_amount(form.get("amount", ""))

    term_raw = form.get("term_date", "").strip()
    if term_raw:
        try:
            deal.term_date = datetime.strptime(term_raw, "%Y-%m-%d").date()
        except ValueError:
            pass
    else:
        deal.term_date = None

    status = form.get("status", "")
    if status in DEAL_STATUSES:
        deal.status = status

    deal.analyst_name = form.get("analyst_name", "").strip()[:300]
    deal.analyst_phone = form.get("analyst_phone", "").strip()[:100]
    deal.comment = form.get("comment", "").strip()


@app.route("/deals")
@login_required
def deals_page():
    status = request.args.get("status", "").strip()
    edit_id = request.args.get("edit", type=int)
    new_mode = request.args.get("new") == "1"
    company_id = request.args.get("company_id", type=int)

    query = Deal.query
    if status:
        query = query.filter_by(status=status)

    deals = query.order_by(Deal.updated_at.desc()).all()

    editing = db.session.get(Deal, edit_id) if edit_id else None
    companies = Company.query.order_by(Company.name).all()
    products = Product.query.filter_by(is_active=True).order_by(Product.name).all()

    status_logs = (
        DealStatusLog.query.order_by(DealStatusLog.created_at.desc()).limit(50).all()
    )

    return render_template(
        "deals.html",
        deals=deals,
        status=status,
        editing=editing,
        new_mode=new_mode or (editing is None and company_id is not None),
        pre_company_id=company_id,
        companies=companies,
        products=products,
        status_logs=status_logs,
    )


@app.route("/deals/new", methods=["POST"])
@login_required
def deal_new():
    company_id = request.form.get("company_id", type=int)
    if not company_id or not db.session.get(Company, company_id):
        flash("Выберите компанию для заявки.", "danger")
        return redirect(url_for("deals_page", new=1))

    deal = Deal(company_id=company_id)
    _fill_deal(deal, request.form)
    db.session.add(deal)
    db.session.flush()

    _log_deal_status(deal, "", deal.status)
    db.session.commit()

    flash("Заявка создана.", "success")
    return redirect(url_for("deals_page"))


@app.route("/deal/<int:deal_id>/edit", methods=["POST"])
@login_required
def deal_edit(deal_id):
    deal = db.get_or_404(Deal, deal_id)
    old_status = deal.status

    _fill_deal(deal, request.form)
    db.session.flush()

    _log_deal_status(deal, old_status, deal.status)
    db.session.commit()

    flash("Заявка сохранена.", "success")
    return redirect(url_for("deals_page"))


@app.route("/deal/<int:deal_id>/status", methods=["POST"])
@login_required
def deal_status(deal_id):
    deal = db.get_or_404(Deal, deal_id)
    status = request.form.get("status", "")

    if status in DEAL_STATUSES and status != deal.status:
        old = deal.status
        deal.status = status
        db.session.flush()
        _log_deal_status(deal, old, status)
        db.session.commit()

    return redirect(request.referrer or url_for("deals_page"))


@app.route("/deal/<int:deal_id>/delete", methods=["POST"])
@login_required
def deal_delete(deal_id):
    deal = db.get_or_404(Deal, deal_id)
    db.session.delete(deal)
    db.session.commit()

    flash("Заявка удалена.", "success")
    return redirect(request.referrer or url_for("deals_page"))


# ============================================================
# ПИСЬМА / OUTLOOK
# ============================================================

def _get_company_contact(company_id, contact_id):
    company = db.get_or_404(Company, company_id)
    contact = db.get_or_404(Contact, contact_id)
    if contact.company_id != company.id:
        abort(404)
    return company, contact


def _get_template(template_id=None):
    if template_id:
        template = db.session.get(MailTemplate, template_id)
        if template:
            return template
    return (
        MailTemplate.query.filter_by(is_default=True).first()
        or MailTemplate.query.first()
    )


def _log_email(company, contact, subject, body, status_note=""):
    db.session.add(Activity(
        company_id=company.id,
        contact_id=contact.id,
        user_id=current_user.id,
        type="email",
        subject=subject,
        body=body + (f"\n\n[Статус: {status_note}]" if status_note else ""),
    ))
    db.session.commit()


@app.route("/company/<int:company_id>/contact/<int:contact_id>/email")
@login_required
def email_open(company_id, contact_id):
    company, contact = _get_company_contact(company_id, contact_id)

    if not contact.email:
        flash("У контакта не заполнен e-mail.", "warning")
        return redirect(url_for("company_detail", company_id=company.id, tab="contacts"))

    template = _get_template(request.args.get("template_id", type=int))
    templates = MailTemplate.query.order_by(MailTemplate.is_default.desc(), MailTemplate.name).all()

    if not template:
        flash("Сначала создайте шаблон письма в разделе «Шаблоны».", "warning")
        return redirect(url_for("templates_page"))

    product_id = request.args.get("product_id", type=int)
    product = db.session.get(Product, product_id) if product_id else None

    subject, body = prepare_email(company, contact, template, product, current_user)

    if request.args.get("auto") == "1":
        ok, message = try_open_outlook(contact.email, subject, body)
        if ok:
            _log_email(company, contact, subject, body, "Черновик открыт в Outlook")
            flash(message, "success")
            return redirect(url_for("company_detail", company_id=company.id, tab="contacts"))
        flash(message, "warning")

    return render_template(
        "compose.html",
        company=company,
        contact=contact,
        subject=subject,
        body=body,
        template=template,
        templates=templates,
        mailto=mailto_link(contact.email, subject, body),
        greeting=build_greeting(contact.full_name),
    )


@app.route("/company/<int:company_id>/contact/<int:contact_id>/email/send", methods=["POST"])
@login_required
def email_send(company_id, contact_id):
    company, contact = _get_company_contact(company_id, contact_id)

    subject = request.form.get("subject", "").strip()
    body = request.form.get("body", "")
    template_id = request.form.get("template_id", type=int)

    template = _get_template(template_id)
    templates = MailTemplate.query.order_by(MailTemplate.is_default.desc(), MailTemplate.name).all()

    if not contact.email:
        flash("У контакта не заполнен e-mail.", "danger")
        return redirect(url_for("company_detail", company_id=company.id, tab="contacts"))

    ok, message = try_open_outlook(contact.email, subject, body)

    if ok:
        _log_email(company, contact, subject, body, "Черновик открыт в Outlook")
        flash(message, "success")
        return redirect(url_for("company_detail", company_id=company.id, tab="contacts"))

    flash(message, "warning")
    return render_template(
        "compose.html",
        company=company,
        contact=contact,
        subject=subject,
        body=body,
        template=template,
        templates=templates,
        mailto=mailto_link(contact.email, subject, body),
        greeting=build_greeting(contact.full_name),
    )


@app.route("/company/<int:company_id>/contact/<int:contact_id>/email/mailto")
@login_required
def email_mailto(company_id, contact_id):
    company, contact = _get_company_contact(company_id, contact_id)

    subject = request.args.get("subject", "")
    body = request.args.get("body", "")

    _log_email(company, contact, subject, body, "Открыт почтовый клиент (mailto)")
    return redirect(mailto_link(contact.email, subject, body))


# ============================================================
# ИСТОРИЯ И ЗАДАЧИ
# ============================================================

@app.route("/company/<int:company_id>/activity/add", methods=["POST"])
@login_required
def activity_add(company_id):
    company = db.get_or_404(Company, company_id)

    activity_type = request.form.get("type", "note")
    if activity_type not in ACTIVITY_TYPES:
        activity_type = "note"

    body = request.form.get("body", "").strip()
    if not body:
        flash("Введите текст заметки.", "danger")
        return redirect(url_for("company_detail", company_id=company.id, tab="history"))

    db.session.add(Activity(
        company_id=company.id,
        user_id=current_user.id,
        type=activity_type,
        subject=request.form.get("subject", "").strip()[:500],
        body=body,
    ))
    db.session.commit()

    flash("Запись добавлена в историю.", "success")
    return redirect(url_for("company_detail", company_id=company.id, tab="history"))


@app.route("/company/<int:company_id>/task/add", methods=["POST"])
@login_required
def task_add(company_id):
    company = db.get_or_404(Company, company_id)

    title = request.form.get("title", "").strip()
    if not title:
        flash("Введите название задачи.", "danger")
        return redirect(url_for("company_detail", company_id=company.id, tab="tasks"))

    due_raw = request.form.get("due_date", "").strip()
    due_date = None
    if due_raw:
        try:
            due_date = datetime.strptime(due_raw, "%Y-%m-%d").date()
        except ValueError:
            pass

    db.session.add(Task(
        company_id=company.id,
        user_id=current_user.id,
        title=title,
        due_date=due_date,
    ))
    db.session.commit()

    flash("Задача создана.", "success")
    return redirect(url_for("company_detail", company_id=company.id, tab="tasks"))


@app.route("/task/<int:task_id>/toggle", methods=["POST"])
@login_required
def task_toggle(task_id):
    task = db.get_or_404(Task, task_id)
    task.status = "done" if task.status == "open" else "open"
    db.session.commit()
    return redirect(request.referrer or url_for("tasks_page"))


@app.route("/task/<int:task_id>/delete", methods=["POST"])
@login_required
def task_delete(task_id):
    task = db.get_or_404(Task, task_id)
    db.session.delete(task)
    db.session.commit()
    flash("Задача удалена.", "success")
    return redirect(request.referrer or url_for("tasks_page"))


@app.route("/tasks")
@login_required
def tasks_page():
    status = request.args.get("status", "open")
    query = Task.query

    if status in ("open", "done"):
        query = query.filter_by(status=status)

    tasks = query.order_by(Task.status, Task.due_date.is_(None), Task.due_date).all()
    return render_template("tasks.html", tasks=tasks, status=status)


# ============================================================
# ПРОДУКТЫ И СПРАВОЧНИК ОКВЭД
# ============================================================

@app.route("/products")
@login_required
def products_page():
    products = Product.query.order_by(Product.name).all()
    versions = ReferenceVersion.query.order_by(ReferenceVersion.id.desc()).limit(10).all()
    active_version = ReferenceVersion.query.filter_by(is_active=True).first()

    stats = {}
    for product in products:
        stats[product.id] = {
            "rules": OkvedRule.query.filter_by(
                product_id=product.id,
                version_id=active_version.id if active_version else -1,
            ).count(),
            "allowed": OkvedCheck.query.filter_by(
                product_id=product.id, result="allowed"
            ).count(),
            "conditional": OkvedCheck.query.filter_by(
                product_id=product.id, result="conditional"
            ).count(),
        }

    return render_template(
        "products.html",
        products=products,
        versions=versions,
        active_version=active_version,
        stats=stats,
    )


@app.route("/products/upload", methods=["POST"])
@login_required
def product_upload():
    file = request.files.get("file")
    if not file or not file.filename:
        flash("Выберите Excel-файл со справочником ОКВЭД.", "danger")
        return redirect(url_for("products_page"))

    filename = secure_filename(file.filename) or "okved_reference.xlsx"
    saved_name = f"{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}_{filename}"
    saved_path = REFERENCES_DIR / saved_name
    file.save(saved_path)

    try:
        result = okved_service.import_reference(str(saved_path), file.filename)
    except Exception as exc:
        db.session.rollback()
        flash(f"Не удалось прочитать справочник: {exc}", "danger")
        return redirect(url_for("products_page"))

    companies_count = okved_service.recheck_all_companies()

    message = (
        f"Загружено: продуктов — {result['products']}, правил ОКВЭД — {result['rules']}. "
        f"Перепроверено компаний — {companies_count}."
    )
    if result.get("skipped"):
        message += f" Листы без таблицы ОКВЭД пропущены: {', '.join(result['skipped'])}."

    flash(message, "success")
    return redirect(url_for("products_page"))


@app.route("/product/<int:product_id>")
@login_required
def product_detail(product_id):
    product = db.get_or_404(Product, product_id)
    version = ReferenceVersion.query.filter_by(is_active=True).first()

    rules = []
    if version:
        rules = (
            OkvedRule.query.filter_by(product_id=product.id, version_id=version.id)
            .order_by(OkvedRule.code)
            .all()
        )

    checks = (
        OkvedCheck.query.filter_by(product_id=product.id)
        .order_by(OkvedCheck.result)
        .all()
    )

    return render_template(
        "product_detail.html",
        product=product,
        rules=rules,
        checks=checks,
        version=version,
    )


@app.route("/product/<int:product_id>/toggle", methods=["POST"])
@login_required
def product_toggle(product_id):
    product = db.get_or_404(Product, product_id)
    product.is_active = not product.is_active
    db.session.commit()
    flash("Статус продукта обновлён.", "success")
    return redirect(url_for("products_page"))


@app.route("/products/recheck", methods=["POST"])
@login_required
def products_recheck():
    count = okved_service.recheck_all_companies()
    flash(f"Перепроверено компаний: {count}.", "success")
    return redirect(url_for("products_page"))


# ============================================================
# ОТЧЁТЫ И ВЫГРУЗКИ
# ============================================================

HEADER_FILL = PatternFill("solid", fgColor="0066CC")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _xlsx_response(headers, rows, filename):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчет"

    for col, name in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=value)

    for col_idx in range(1, len(headers) + 1):
        sheet.column_dimensions[
            sheet.cell(row=1, column=col_idx).column_letter
        ].width = 22

    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/reports")
@login_required
def reports_page():
    pipeline = []
    for key, label in DEAL_STATUSES.items():
        deals = Deal.query.filter_by(status=key).all()
        pipeline.append({
            "key": key,
            "label": label,
            "count": len(deals),
            "amount": sum(d.amount or 0 for d in deals),
        })

    total_active = sum(
        d.amount or 0 for d in Deal.query.filter(
            Deal.status.notin_(["issued", "rejected"])
        ).all()
    )

    return render_template(
        "reports.html",
        pipeline=pipeline,
        total_active=total_active,
        companies_count=Company.query.count(),
        deals_count=Deal.query.count(),
    )


@app.route("/export/companies")
@login_required
def export_companies():
    headers = [
        "ИНН", "Наименование", "Краткое наименование", "Статус работы",
        "Основной ОКВЭД", "Вид деятельности", "Руководитель",
        "Телефон", "E-mail", "Сайт",
        "Доступно продуктов", "С условиями", "Дата обновления",
    ]

    rows = []
    for company in Company.query.order_by(Company.name).all():
        rows.append([
            company.inn,
            company.name,
            company.short_name,
            company.status_label,
            company.okved_main_code,
            company.okved_main_title,
            company.director,
            company.phone,
            company.email,
            company.site,
            company.allowed_products_count,
            company.conditional_products_count,
            company.checko_updated_at.strftime("%d.%m.%Y %H:%M") if company.checko_updated_at else "",
        ])

    return _xlsx_response(headers, rows, "companies.xlsx")


@app.route("/export/contacts")
@login_required
def export_contacts():
    headers = ["Компания", "ИНН", "ФИО", "Должность", "Телефон", "E-mail", "Основной"]

    rows = []
    for contact in Contact.query.join(Company).order_by(Company.name).all():
        rows.append([
            contact.company.display_name,
            contact.company.inn,
            contact.full_name,
            contact.position,
            contact.phone,
            contact.email,
            "да" if contact.is_primary else "",
        ])

    return _xlsx_response(headers, rows, "contacts.xlsx")


@app.route("/export/deals")
@login_required
def export_deals():
    headers = [
        "Компания", "ИНН", "Продукт", "Сумма", "Срок",
        "Статус", "Аналитик", "Телефон аналитика", "Комментарий",
        "Создана", "Обновлена",
    ]

    rows = []
    for deal in Deal.query.join(Company).order_by(Deal.updated_at.desc()).all():
        rows.append([
            deal.company.display_name,
            deal.company.inn,
            deal.product.name if deal.product else "",
            deal.amount or 0,
            deal.term_date.strftime("%d.%m.%Y") if deal.term_date else "",
            deal.status_label,
            deal.analyst_name,
            deal.analyst_phone,
            deal.comment,
            deal.created_at.strftime("%d.%m.%Y"),
            deal.updated_at.strftime("%d.%m.%Y %H:%M"),
        ])

    return _xlsx_response(headers, rows, "deals.xlsx")


@app.route("/export/tasks")
@login_required
def export_tasks():
    headers = ["Компания", "ИНН", "Задача", "Срок", "Статус"]

    rows = []
    for task in Task.query.join(Company).order_by(Task.due_date).all():
        rows.append([
            task.company.display_name,
            task.company.inn,
            task.title,
            task.due_date.strftime("%d.%m.%Y") if task.due_date else "",
            "Выполнена" if task.status == "done" else ("Просрочена" if task.is_overdue else "В работе"),
        ])

    return _xlsx_response(headers, rows, "tasks.xlsx")


@app.route("/export/products")
@login_required
def export_products():
    products = Product.query.filter_by(is_active=True).order_by(Product.name).all()
    headers = ["Компания", "ИНН", "Основной ОКВЭД"] + [p.name for p in products]

    labels = {
        "allowed": "Доступно",
        "conditional": "С условиями",
        "forbidden": "Недоступно",
        "unknown": "Нет данных",
    }

    rows = []
    for company in Company.query.order_by(Company.name).all():
        results = {c.product_id: c.result for c in company.checks}
        row = [company.display_name, company.inn, company.okved_main_code]
        for product in products:
            row.append(labels.get(results.get(product.id), ""))
        rows.append(row)

    return _xlsx_response(headers, rows, "products_matrix.xlsx")


@app.route("/export/deal_log")
@login_required
def export_deal_log():
    headers = ["Дата", "Компания", "ИНН", "Продукт", "Старый статус", "Новый статус", "Кто изменил"]

    rows = []
    for log in DealStatusLog.query.order_by(DealStatusLog.created_at.desc()).all():
        rows.append([
            log.created_at.strftime("%d.%m.%Y %H:%M"),
            log.company.display_name,
            log.company.inn,
            log.deal.product.name if log.deal and log.deal.product else "",
            DEAL_STATUSES.get(log.old_status, log.old_status or "—"),
            DEAL_STATUSES.get(log.new_status, log.new_status),
            log.user.full_name or (log.user.username if log.user else ""),
        ])

    return _xlsx_response(headers, rows, "deal_status_log.xlsx")


# ============================================================
# ШАБЛОНЫ ПИСЕМ
# ============================================================

@app.route("/templates", methods=["GET", "POST"])
@login_required
def templates_page():
    if request.method == "POST":
        template_id = request.form.get("template_id", type=int)

        if template_id:
            template = db.get_or_404(MailTemplate, template_id)
        else:
            template = MailTemplate()
            db.session.add(template)

        template.name = request.form.get("name", "").strip()[:200]
        template.subject = request.form.get("subject", "").strip()[:500]
        template.body = request.form.get("body", "")
        template.is_default = request.form.get("is_default") == "on"

        if not template.name:
            flash("Введите название шаблона.", "danger")
            db.session.rollback()
            return redirect(url_for("templates_page"))

        if template.is_default:
            MailTemplate.query.filter(MailTemplate.id != template.id).update({"is_default": False})

        db.session.commit()
        flash("Шаблон сохранён.", "success")
        return redirect(url_for("templates_page"))

    templates = MailTemplate.query.order_by(MailTemplate.is_default.desc(), MailTemplate.name).all()
    return render_template("templates_page.html", templates=templates, selected=None)


@app.route("/templates/<int:template_id>/edit")
@login_required
def template_edit(template_id):
    template = db.get_or_404(MailTemplate, template_id)
    templates = MailTemplate.query.order_by(MailTemplate.is_default.desc(), MailTemplate.name).all()
    return render_template("templates_page.html", templates=templates, selected=template)


@app.route("/templates/<int:template_id>/delete", methods=["POST"])
@login_required
def template_delete(template_id):
    template = db.get_or_404(MailTemplate, template_id)

    if MailTemplate.query.count() == 1:
        flash("Нельзя удалить единственный шаблон.", "warning")
        return redirect(url_for("templates_page"))

    db.session.delete(template)
    db.session.commit()
    flash("Шаблон удалён.", "success")
    return redirect(url_for("templates_page"))


# ============================================================
# СТАТУС ФОНОВОЙ ЗАДАЧИ
# ============================================================

@app.route("/jobs/<job_id>")
@login_required
def job_status(job_id):
    job = _job_get(job_id)
    if not job:
        return jsonify({"status": "error", "error": "Задача не найдена."}), 404

    job.pop("updated_at", None)
    return jsonify(job)


# ============================================================
# ОШИБКИ
# ============================================================

@app.errorhandler(404)
def not_found(error):
    return render_template("base.html", error_page="404. Страница не найдена."), 404


@app.errorhandler(500)
def server_error(error):
    db.session.rollback()
    return render_template("base.html", error_page="500. Внутренняя ошибка сервера."), 500


# ============================================================
# АВТООТКРЫТИЕ БРАУЗЕРА С CRM
# ============================================================

def open_browser_delayed(delay=1.5):
    def _worker():
        time.sleep(delay)
        try:
            webbrowser.open(Config.APP_URL)
        except Exception as exc:
            print(f"Не удалось открыть браузер: {exc}")

    threading.Thread(target=_worker, daemon=True).start()


# ============================================================
# ЗАПУСК
# ============================================================

if __name__ == "__main__":
    init_database()

    print("=" * 60)
    print("CRM клиентского менеджера банка")
    print(f"Адрес: {Config.APP_URL}")
    print("Логин: manager   Пароль: manager")
    if Config.AUTO_OPEN_BROWSER:
        print("Браузер откроется автоматически…")
    print("=" * 60)

    if Config.AUTO_OPEN_BROWSER:
        open_browser_delayed(1.5)

    app.run(host="127.0.0.1", port=5000, debug=True, use_reloader=False)